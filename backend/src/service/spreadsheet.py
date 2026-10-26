import csv
import datetime
import io
import threading
import uuid

from collections.abc import Sequence

import openpyxl

import schema.rest

CacheEntry = tuple[dict[str, tuple[list[str], list[list[str]], int]], datetime.datetime]
_cache: dict[uuid.UUID, CacheEntry] = {}
_cache_lock = threading.Lock()
_CACHE_TTL = datetime.timedelta(minutes=30)


def _cleanup(task_id: uuid.UUID) -> None:
    with _cache_lock:
        _cache.pop(task_id, None)


def _take_first_nonempty_as_header(
    rows: Sequence[Sequence],
) -> tuple[list[str], list[list[str]], int]:
    """Skip leading blank rows; use the first non-empty row as the column headers.

    Returns (columns, data_rows, header_row) where header_row is 1-based
    in the original spreadsheet (first non-blank row)."""
    if not rows:
        return [], [], 0

    # skip leading rows that are entirely blank
    start = 0
    for i, row in enumerate(rows):
        if any(c is not None and str(c).strip() for c in row):
            start = i
            break
    else:
        return [], [], 0

    columns = [str(c) if c is not None else "" for c in rows[start]]
    data_rows = rows[start + 1 :]
    data = [[str(c) if c is not None else "" for c in row] for row in data_rows]
    header_row = start + 1  # convert 0-based index to 1-based sheet row
    return columns, data, header_row


def _parse_xlsx(file_bytes: bytes) -> dict[str, tuple[list[str], list[list[str]], int]]:
    wb: openpyxl.Workbook | None = None
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(file_bytes), read_only=True, data_only=True
        )
        result: dict[str, tuple[list[str], list[list[str]], int]] = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                result[sheet_name] = ([], [], 0)
                continue
            columns, data, header_row = _take_first_nonempty_as_header(rows)
            result[sheet_name] = (columns, data, header_row)
        return result
    finally:
        if wb is not None:
            wb.close()


def _parse_csv(file_bytes: bytes) -> dict[str, tuple[list[str], list[list[str]], int]]:
    text = file_bytes.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return {"Sheet1": ([], [], 0)}
    columns, data, header_row = _take_first_nonempty_as_header(rows)
    return {"Sheet1": (columns, data, header_row)}


def parse_workbook(
    file_bytes: bytes, filename: str
) -> dict[str, tuple[list[str], list[list[str]], int]]:
    name_lower = filename.lower()
    if name_lower.endswith(".csv"):
        return _parse_csv(file_bytes)
    if name_lower.endswith(".xlsx"):
        return _parse_xlsx(file_bytes)
    raise ValueError("Only .xlsx and .csv files are supported")


def store_workbook(
    workbook: dict[str, tuple[list[str], list[list[str]], int]],
) -> uuid.UUID:
    task_id = uuid.uuid4()
    expiry = datetime.datetime.now(datetime.timezone.utc) + _CACHE_TTL
    with _cache_lock:
        _cache[task_id] = (workbook, expiry)
    threading.Timer(_CACHE_TTL.total_seconds(), _cleanup, args=(task_id,)).start()
    return task_id


def get_workbook(
    task_id: uuid.UUID,
) -> dict[str, tuple[list[str], list[list[str]], int]] | None:
    with _cache_lock:
        entry = _cache.get(task_id)
        if entry is None:
            return None
        workbook, expiry = entry
        if datetime.datetime.now(datetime.timezone.utc) > expiry:
            del _cache[task_id]
            return None
        return workbook


def map_rows(
    workbook: dict[str, tuple[list[str], list[list[str]], int]],
    sheet_name: str,
    title_column: str | None,
    name_column: str | None,
    raw_phone_column: str | None,
    email_column: str | None,
) -> list[schema.rest.ParticipantCreate]:
    if sheet_name not in workbook:
        raise ValueError(f"Sheet '{sheet_name}' not found in workbook")
    columns, rows, _ = workbook[sheet_name]
    idx_by_field: dict[str, int] = {}
    if title_column is not None and title_column in columns:
        idx_by_field["title"] = columns.index(title_column)
    if name_column is not None and name_column in columns:
        idx_by_field["name"] = columns.index(name_column)
    if raw_phone_column is not None and raw_phone_column in columns:
        idx_by_field["raw_phone"] = columns.index(raw_phone_column)
    if email_column is not None and email_column in columns:
        idx_by_field["email"] = columns.index(email_column)

    def _get(row: list[str], idx: int) -> str:
        return row[idx] if idx < len(row) else ""

    result: list[schema.rest.ParticipantCreate] = []
    for row in rows:
        title = (
            _get(row, idx_by_field["title"]).strip()
            if "title" in idx_by_field
            else None
        )
        name = (
            _get(row, idx_by_field["name"]).strip() if "name" in idx_by_field else None
        )
        raw_phone = (
            _get(row, idx_by_field["raw_phone"]).strip()
            if "raw_phone" in idx_by_field
            else None
        )
        email = (
            _get(row, idx_by_field["email"]).strip()
            if "email" in idx_by_field
            else None
        )
        if not name:
            continue
        result.append(
            schema.rest.ParticipantCreate(
                title=title, name=name, email=email, raw_phone=raw_phone
            )
        )
    return result
